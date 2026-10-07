"""Printable claim cover and safe, explicit filename templates."""
from datetime import datetime
from decimal import Decimal, DecimalException, ROUND_HALF_UP
import re
import math

from .amount_utils import parse_amount

_DIGITS = '零壹贰叁肆伍陆柒捌玖'


def _group(number):
    result, pending = '', False
    for power, unit in [(1000, '仟'), (100, '佰'), (10, '拾'), (1, '')]:
        digit, number = divmod(number, power)
        if digit:
            if pending:
                result += '零'
            result += _DIGITS[digit] + unit
            pending = False
        elif result and number:
            pending = True
    return result


def rmb_upper(value):
    try:
        amount = Decimal(str(value))
    except DecimalException:
        raise ValueError("大写金额格式无效") from None
    if not amount.is_finite() or abs(amount) >= Decimal('10000000000000000'):
        raise ValueError('大写金额超出支持范围')
    amount = amount.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    if abs(amount) >= Decimal('10000000000000000'):
        raise ValueError('大写金额超出支持范围')
    negative = amount < 0
    cents = int(abs(amount) * 100)
    integer, fraction = divmod(cents, 100)
    groups = []
    while integer:
        integer, part = divmod(integer, 10000)
        groups.append(part)
    result, zero = '', False
    for index in range(len(groups) - 1, -1, -1):
        part = groups[index]
        if not part:
            if result:
                zero = True
            continue
        if result and (zero or part < 1000):
            result += '零'
        result += _group(part) + ['', '万', '亿', '万亿'][index]
        zero = False
    result = (result or '零') + '元'
    jiao, fen = divmod(fraction, 10)
    if not fraction:
        result += '整'
    else:
        if jiao:
            result += _DIGITS[jiao] + '角'
        elif cents >= 100:
            result += '零'
        if fen:
            result += _DIGITS[fen] + '分'
    return ('负' if negative else '') + result


def claim_filename(template, claim, now=None):
    now = now or datetime.now()
    values = {'YYYYMM': now.strftime('%Y%m'), '部门': claim.get('department') or '未填部门',
              '姓名': claim.get('applicant_name') or '未填姓名',
              '事由': claim.get('reason_detail') or claim.get('reason_category') or claim['name']}
    if template is not None and not isinstance(template, str):
        raise ValueError('文件名模板必须是文本')
    template = template or 'reimbursement.xlsx'
    tokens = re.findall(r'\{([^{}]+)\}', template)
    if any(token not in values for token in tokens):
        raise ValueError('文件名模板仅支持 {YYYYMM}、{部门}、{姓名}、{事由}')
    remainder = re.sub(r'\{[^{}]+\}', '', template)
    if '{' in remainder or '}' in remainder:
        raise ValueError('文件名模板括号不完整')
    name = re.sub(r'\{([^{}]+)\}', lambda match: str(values[match[1]]), template)
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f\x7f]', '_', name).strip().rstrip('. ')
    stem = name[:-5] if name.lower().endswith('.xlsx') else name
    stem = stem.rstrip('. ')[:110].rstrip('. ') or 'reimbursement'
    if re.fullmatch(r'(?i)(CON|PRN|AUX|NUL|COM[1-9¹²³]|LPT[1-9¹²³])', stem.split('.', 1)[0].rstrip('. ')):
        stem = '_' + stem
    return stem + '.xlsx'


def add_claim_cover(wb, rows, claim):
    from openpyxl.styles import Alignment, Font, Border, Side
    from openpyxl.worksheet.page import PageMargins
    ws = wb.create_sheet('费用报销单', 0)
    totals = {}
    for row in rows:
        currency = str(row.get('currency') or 'CNY').strip().upper()
        currency = 'CNY' if currency in {'RMB', '人民币'} else currency
        totals[currency] = totals.get(currency, Decimal('0')) + parse_amount(row['total_amount'])
    total = totals.get('CNY', Decimal('0'))
    amount_text = format(total, '.2f') if set(totals) <= {'CNY'} else '\n'.join(
        f'{currency} {amount:.2f}' for currency, amount in sorted(totals.items()))
    upper_text = rmb_upper(total) if set(totals) <= {'CNY'} else '含外币，人民币大写不适用（未折算）'
    entries = [('费用报销单', ''), ('报销组', claim['name']),
               ('报销人', claim.get('applicant_name') or ''), ('部门', claim.get('department') or ''),
               ('事由分类', claim.get('reason_category') or ''), ('事由说明', claim.get('reason_detail') or ''),
               ('票据数量', len(rows)), ('报销金额', amount_text),
               ('金额大写', upper_text), ('申请人签字', ''), ('审批人签字', ''), ('财务复核', ''),
               ('报销期间', ' ~ '.join(str(claim.get(key) or '') for key in ('period_start', 'period_end'))
                if claim.get('period_start') or claim.get('period_end') else ''),
               ('生成日期', datetime.now().strftime('%Y-%m-%d'))]
    for index, (label, value) in enumerate(entries, 1):
        ws.cell(index, 1, label)
        ws.merge_cells(start_row=index, start_column=2, end_row=index, end_column=6)
        cell = ws.cell(index, 2, str(value))
        cell.data_type = 's'
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.font = Font(name='微软雅黑', size=12)
        line_count = max(1, sum(max(1, math.ceil(sum(2 if ord(char) > 127 else 1 for char in line) / 55))
                                 for line in str(value).split('\n')))
        ws.row_dimensions[index].height = max(72 if index == 6 else 36, line_count * 18)
        ws.cell(index, 1).font = Font(name='微软雅黑', bold=True, size=12)
    edge = Side(style='thin', color='808080')
    for row in ws.iter_rows(min_row=2, max_row=14, min_col=1, max_col=6):
        for cell in row:
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
    ws.merge_cells('A1:F1')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws['A1'].font = Font(name='微软雅黑', bold=True, size=18)
    ws.column_dimensions['A'].width = 18
    for column in 'BCDEF':
        ws.column_dimensions[column].width = 12
    ws.sheet_view.showGridLines = False
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=.4, right=.4, top=.4, bottom=.4)
    ws.print_area = 'A1:F14'
