"""Exercise startup/export boundaries in fresh Python processes."""

import os
from pathlib import Path
import subprocess
import sys
import tempfile
import textwrap
import unittest


PROJECT_ROOT = Path(__file__).resolve().parents[1]
BLOCK_SPREADSHEETS = """
import importlib.abc
import sys
class BlockSpreadsheets(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname == 'openpyxl' or fullname.startswith('openpyxl.'):
            raise AssertionError('spreadsheet dependency loaded before export: ' + fullname)
sys.meta_path.insert(0, BlockSpreadsheets())
"""


class StartupDependencyTests(unittest.TestCase):
    def run_script(self, source, *args):
        env = os.environ.copy()
        env.update(QT_QPA_PLATFORM="offscreen", INVOICE_HUB_TEST_MODE="1", PYTHONUTF8="1")
        with tempfile.TemporaryDirectory(prefix="invoice-hub-dependencies-") as td:
            env["INVOICE_HUB_RUNTIME_DIR"] = str(Path(td) / "runtime")
            result = subprocess.run(
                [sys.executable, "-c", textwrap.dedent(source), *args],
                cwd=PROJECT_ROOT, env=env, capture_output=True, text=True,
                encoding="utf-8", timeout=45,
            )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        return result.stdout

    def test_gui_cli_and_service_imports_do_not_require_spreadsheets(self):
        for module in (
            "scripts.invoice_fetch.__main__",
            "scripts.invoice_fetch.services",
            "scripts.invoice_fetch.gui.app",
            "scripts.invoice_fetch.gui.startup_lifecycle",
            "scripts.invoice_fetch.gui.startup_probe",
            "scripts.invoice_fetch.gui.business_pages_baseline",
            "scripts.invoice_fetch.claim_export",
        ):
            with self.subTest(module=module):
                self.run_script(BLOCK_SPREADSHEETS + "\nimport importlib\nimportlib.import_module(sys.argv[1])", module)

    def test_material_and_naming_preflight_work_without_spreadsheets(self):
        self.run_script(BLOCK_SPREADSHEETS + textwrap.dedent("""
            from pathlib import Path
            import tempfile
            from scripts.invoice_fetch.claim_export import summarize_extra_material_issues
            from scripts.invoice_fetch.gui.business_pages_baseline import _export_naming_state
            with tempfile.TemporaryDirectory() as td:
                original = Path(td) / 'synthetic.xml'
                original.write_text('<invoice/>', encoding='utf-8')
                rows = [{'review_status': 'approved', 'invoice_date': '2026/07/03',
                         'attachment_path': str(original), 'missing_extra': 0}]
                assert summarize_extra_material_issues(rows, Path(td)) == {
                    'missing_attachment': 0, 'missing_extra': 0, 'unavailable_extra': 0}
                assert _export_naming_state(rows) == ('1 张使用日期前缀 + 原文件名', 'success')
                rows[0]['invoice_date'] = ''
                assert _export_naming_state(rows) == ('1 张将使用 unknown-date 前缀', 'warning')
        """))

    def test_first_export_loads_spreadsheets_and_preserves_workbook_styles(self):
        self.run_script("""
            import sys, tempfile
            from pathlib import Path
            from scripts.invoice_fetch.excel_export import export_excel
            assert 'openpyxl' not in sys.modules
            with tempfile.TemporaryDirectory() as td:
                destination = Path(td) / 'synthetic.xlsx'
                rows = [{'invoice_number': 'SYNTHETIC', 'total_amount': '1.23',
                         'category': '交通', 'attachment_path': 'synthetic.xml', 'parse_success': 1}]
                export_excel(rows, destination, claim={'name': 'Synthetic'})
                assert 'openpyxl' in sys.modules
                from openpyxl import load_workbook
                for _ in range(2):
                    with destination.open('rb') as stream:
                        book = load_workbook(stream)
                        try:
                            assert book.sheetnames == ['费用报销单', '发票汇总', '分类汇总', '异常待处理']
                            assert book['费用报销单']['B8'].value == '1.23'
                            assert book['发票汇总']['A1'].font.bold
                            assert book['发票汇总']['A1'].fill.fgColor.rgb == '004472C4'
                            assert book['发票汇总']['A2'].border.bottom.style == 'thin'
                            assert book['分类汇总']['C3'].value == 1.23
                            assert book['分类汇总']['C3'].font.bold
                        finally:
                            book.close()
                    export_excel(rows, destination, claim={'name': 'Synthetic'})
        """)

    def test_public_desktop_reaches_first_paint_without_spreadsheets(self):
        output = self.run_script(BLOCK_SPREADSHEETS + textwrap.dedent("""
            import runpy
            from scripts.invoice_fetch.gui.startup_probe import StartupProbeSession
            finish = StartupProbeSession._finish_after_paint
            def observe(self):
                assert self._window._startup_first_paint_seen
                assert not any(m == 'openpyxl' or m.startswith('openpyxl.') for m in sys.modules)
                print('SPREADSHEET_FREE_FIRST_PAINT=1', flush=True)
                finish(self)
            StartupProbeSession._finish_after_paint = observe
            sys.argv = ['invoice_fetch', 'desktop', '--startup-probe']
            runpy.run_module('scripts.invoice_fetch', run_name='__main__')
        """))
        self.assertIn("SPREADSHEET_FREE_FIRST_PAINT=1", output)
        self.assertIn("QT_PAINT_EVENT_COMPLETED=1", output)
        self.assertIn("PROBE_CONTRACT=main_window_first_paint_v1", output)


if __name__ == "__main__":
    unittest.main()
