import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from scripts.invoice_fetch.claim_cover import rmb_upper, claim_filename
from scripts.invoice_fetch.excel_export import export_excel


class ClaimCoverTests(unittest.TestCase):
    def test_upper_amounts(self):
        cases = {'0': '零元整', '10': '壹拾元整', '10001': '壹万零壹元整',
                 '100000001': '壹亿零壹元整', '1000000000000': '壹万亿元整', '100010001': '壹亿零壹万零壹元整',
                 '1.01': '壹元零壹分', '0.01': '零元壹分', '0.10': '零元壹角',
                 '-13.02': '负壹拾叁元零贰分', '1.005': '壹元零壹分'}
        for value, expected in cases.items():
            with self.subTest(value=value):
                self.assertEqual(rmb_upper(value), expected)
        for value in ('NaN', 'Infinity', '10000000000000000'):
            with self.assertRaises(ValueError):
                rmb_upper(value)

    def test_safe_template(self):
        claim = {'name': 'Group', 'department': '../研发', 'applicant_name': '李某', 'reason_detail': '验收/项目'}
        name = claim_filename('{YYYYMM}_{部门}_{姓名}_{事由}.xlsx', claim, datetime(2026, 10, 7))
        self.assertEqual(name, '202610_.._研发_李某_验收_项目.xlsx')
        self.assertEqual(claim_filename('CON.xlsx', claim), '_CON.xlsx')
        self.assertEqual(claim_filename('CON .foo.xlsx', claim), '_CON .foo.xlsx')
        self.assertEqual(claim_filename('COM¹.xlsx', claim), '_COM¹.xlsx')
        self.assertEqual(claim_filename('{事由}', {'name': 'Group', 'reason_detail': '项目{一期}'}), '项目{一期}.xlsx')
        self.assertEqual(claim_filename(None, claim), 'reimbursement.xlsx')
        self.assertLessEqual(len(claim_filename('长' * 500, claim)), 115)
        with self.assertRaises(ValueError):
            claim_filename('{bad}', claim)

    def test_cover_print_settings_static_values_and_text_safety(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'claim.xlsx'
            rows = [{'invoice_number': '1', 'total_amount': '0.10'},
                    {'invoice_number': '2', 'total_amount': '0.20'}]
            export_excel(rows, path, claim={'name': 'Group', 'reason_detail': '=SUM(A1:A2)'})
            wb = load_workbook(BytesIO(path.read_bytes()), data_only=True)
            try:
                self.assertEqual(wb.sheetnames[0], '费用报销单')
                sheet = wb['费用报销单']
                self.assertEqual(sheet['B8'].value, '0.30')
                self.assertEqual(sheet['B9'].value, '零元叁角')
                self.assertEqual(sheet['B6'].value, '=SUM(A1:A2)')
                self.assertEqual(sheet['B6'].data_type, 's')
                self.assertEqual(str(sheet.page_setup.paperSize), sheet.PAPERSIZE_A4)
                self.assertEqual(sheet.page_setup.fitToWidth, 1)
                self.assertIn('发票汇总', wb.sheetnames)
            finally:
                wb.close()

    def test_foreign_amounts_are_not_summed_or_labeled_as_rmb(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'currency.xlsx'
            export_excel([{'total_amount': '10', 'currency': 'USD'},
                          {'total_amount': '20', 'currency': 'CNY'}], path, claim={'name': 'Group'})
            wb = load_workbook(BytesIO(path.read_bytes()), data_only=True)
            try:
                self.assertEqual(wb['费用报销单']['B8'].value, 'CNY 20.00\nUSD 10.00')
                self.assertIn('未折算', wb['费用报销单']['B9'].value)
            finally:
                wb.close()

    def test_package_custom_name_is_recorded_and_cover_contains_group_metadata(self):
        from scripts.invoice_fetch.claim_export import export_claim_package
        from scripts.invoice_fetch.db import InvoiceDB
        import json
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / 'original.pdf').write_bytes(b'synthetic invoice')
            with InvoiceDB(root / 'invoice.db') as db:
                claim = db.create_claim_group('Group', applicant_name='李某', department='研发', reason_detail='项目验收')
                invoice = db.insert_invoice({'invoice_number': 'COVER', 'total_amount': '10',
                                             'review_status': 'approved', 'attachment_path': 'original.pdf'})
                self.assertTrue(db.add_invoice_to_claim(claim, invoice))
                package = export_claim_package(db, claim, root, root, export_root=root / 'exports',
                    reimbursement_config={'filename_template': '{部门}_{姓名}_{事由}.xlsx'})
                manifest = json.loads((package / 'manifest.json').read_text())
                self.assertEqual(manifest['spreadsheet'], '研发_李某_项目验收.xlsx')
                wb = load_workbook(BytesIO((package / manifest['spreadsheet']).read_bytes()), data_only=True)
                try:
                    sheet = wb['费用报销单']
                    self.assertEqual(sheet['B3'].value, '李某')
                    self.assertEqual(sheet['B4'].value, '研发')
                    self.assertEqual(sheet['B6'].value, '项目验收')
                finally:
                    wb.close()

    def test_invalid_template_fails_before_creating_export_folder(self):
        from scripts.invoice_fetch.claim_export import export_claim_package
        from scripts.invoice_fetch.db import InvoiceDB
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / 'original.pdf').write_bytes(b'synthetic invoice')
            with InvoiceDB(root / 'invoice.db') as db:
                claim = db.create_claim_group('Group')
                invoice = db.insert_invoice({'invoice_number': 'INVALID-TEMPLATE', 'total_amount': '10',
                                             'review_status': 'approved', 'attachment_path': 'original.pdf'})
                self.assertTrue(db.add_invoice_to_claim(claim, invoice))
                with self.assertRaises(ValueError):
                    export_claim_package(db, claim, root, root, export_root=root / 'exports',
                                         reimbursement_config={'filename_template': '{unsupported}'})
                self.assertFalse((root / 'exports').exists())

    def test_personalized_filename_is_redacted_in_export_logs(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / '研发_李某_项目验收.xlsx'
            with self.assertLogs('scripts.invoice_fetch.excel_export', level='INFO') as logs:
                export_excel([{'total_amount': '10'}], path, claim={'name': 'Group'})
            text = ' '.join(logs.output)
            self.assertNotIn('李某', text)
            self.assertIn('file#', text)
