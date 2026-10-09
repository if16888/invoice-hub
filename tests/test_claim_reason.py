"""Claim defaults, overrides, migration atomicity and export snapshots."""
import json
from io import BytesIO
import tempfile
import unittest
from pathlib import Path
import sqlite3

from openpyxl import load_workbook

from scripts.invoice_fetch.claim_export import export_claim_package
from scripts.invoice_fetch.db import InvoiceDB
from scripts.invoice_fetch.migrations import check_and_migrate, validate_latest_schema


def add_invoice(db, number='REASON-001', **values):
    return db.insert_invoice({'invoice_number': number, 'total_amount': '10.00',
                              'seller_name': 'Seller', 'review_status': 'approved', **values})


def make_v8(path):
    with InvoiceDB(path) as db:
        invoice = add_invoice(db)
        claim = db.create_claim_group('Legacy')
        assert db.add_invoice_to_claim(claim, invoice)
    with sqlite3.connect(path) as conn:
        conn.execute('ALTER TABLE invoices DROP COLUMN custom_reason')
        for column in ('reason_category', 'reason_detail', 'applicant_name', 'department'):
            conn.execute(f'ALTER TABLE claim_groups DROP COLUMN {column}')
        conn.execute('PRAGMA user_version=8')
    return invoice, claim

class ClaimReasonTests(unittest.TestCase):
    def setUp(self):
        self._tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self._tempdir.cleanup)
        self.tmp_path = Path(self._tempdir.name)

    def test_v8_upgrade_preserves_data_and_unknown_reasons(self):
        tmp_path = self.tmp_path
        path = tmp_path / 'legacy.db'
        invoice, claim = make_v8(path)
        for _ in range(2):
            with InvoiceDB(path) as db:
                validate_latest_schema(db._conn)
                assert db.get_invoice(invoice)['custom_reason'] is None
                assert db.get_claim_group(claim)['reason_category'] == ''
                assert db.get_invoice_reason(invoice) == ''
                assert db.get_invoice_claim_id(invoice) == claim
                assert db.get_invoice(invoice)['total_amount'] == '10.00'
                assert db.count_invoices() == 1


    def test_failed_migration_rolls_back_columns_and_version(self):
        tmp_path = self.tmp_path
        path = tmp_path / 'rollback.db'
        make_v8(path)
        conn = sqlite3.connect(path)
        count = 0

        def authorize(action, arg1, arg2, database, trigger):
            nonlocal count
            if action == sqlite3.SQLITE_ALTER_TABLE:
                count += 1
                if count == 2:
                    return sqlite3.SQLITE_DENY
            return sqlite3.SQLITE_OK

        conn.set_authorizer(authorize)
        try:
            try:
                check_and_migrate(conn)
            except sqlite3.DatabaseError:
                pass
            else:
                raise AssertionError('Expected migration failure')
            conn.set_authorizer(None)
            assert conn.execute('PRAGMA user_version').fetchone()[0] == 8
            assert 'reason_category' not in {r[1] for r in conn.execute('PRAGMA table_info(claim_groups)')}
            check_and_migrate(conn)
            validate_latest_schema(conn)
        finally:
            conn.close()


    def test_inheritance_override_clear_relink_and_reopen(self):
        tmp_path = self.tmp_path
        path = tmp_path / 'reason.db'
        with InvoiceDB(path) as db:
            claim = db.create_claim_group('Project', reason_category='差旅', reason_detail='项目验收')
            other = db.create_claim_group('Other', reason_detail='客户拜访')
            invoice = add_invoice(db)
            assert db.get_invoice_reason(invoice) == ''
            assert db.add_invoice_to_claim(claim, invoice)
            assert db.get_invoice_reason(invoice) == '差旅 · 项目验收'
            assert db.update_invoice_reason(invoice, ' 单独说明 ')
            assert db.update_claim_reason(claim, reason_category='会议', reason_detail='年度会议')
            assert db.get_invoice_reason(invoice) == '单独说明'
            assert db.update_invoice_reason(invoice, '')
            assert db.get_invoice_reason(invoice) == ''
            assert db.update_invoice_reason(invoice, None)
            assert db.get_invoice_reason(invoice) == '会议 · 年度会议'
            assert db.remove_invoice_from_claim(claim, invoice)
            assert db.add_invoice_to_claim(other, invoice)
            assert db.get_claim_invoices(other)[0]['reimbursement_reason'] == '客户拜访'
            assert not db.update_invoice_reason(999999, 'missing')
            assert not db.update_claim_reason(999999, reason_detail='missing')
        with InvoiceDB(path) as db:
            assert db.get_invoice_reason(invoice) == '客户拜访'


    def test_package_excel_and_manifest_snapshot_reason_and_escape_formula(self):
        tmp_path = self.tmp_path
        runtime = tmp_path / 'runtime'
        runtime.mkdir()
        (runtime / 'invoice.pdf').write_bytes(b'synthetic invoice')
        with InvoiceDB(runtime / 'invoice.db') as db:
            claim = db.create_claim_group('Export', reason_detail='项目验收', applicant_name='李某', department='研发')
            invoice = add_invoice(db, attachment_path='invoice.pdf')
            assert db.add_invoice_to_claim(claim, invoice)
            assert db.update_invoice_reason(invoice, '=HYPERLINK("evil")')
            package = export_claim_package(db, claim, tmp_path, runtime, reimbursement_config={}, export_root=tmp_path / "exports")
            manifest = json.loads((package / 'manifest.json').read_text())
            assert manifest['reason_detail'] == '项目验收'
            assert manifest['applicant_name'] == '李某'
            assert manifest['department'] == '研发'
            assert manifest['items'][0]['reimbursement_reason'] == '=HYPERLINK("evil")'
            wb = load_workbook(BytesIO((package / 'reimbursement.xlsx').read_bytes()), data_only=True)
            try:
                ws = wb['发票汇总']
                headers = [c.value for c in ws[1]]
                assert ws.cell(2, headers.index('报销事由') + 1).value == '\'=HYPERLINK("evil")'
            finally:
                wb.close()
            assert not db.update_invoice_reason(invoice, None)
            assert db.get_invoice_reason(invoice) == '=HYPERLINK("evil")'
            assert not db.update_claim_reason(claim, reason_detail='新说明')
            assert db.get_claim_group(claim)['reason_detail'] == '项目验收'
            assert db.get_claim_group(claim)['department'] == '研发'
            assert db.get_claim_group(claim)['applicant_name'] == '李某'
            assert manifest['items'][0]['reimbursement_reason'] == '=HYPERLINK("evil")'


    def test_backup_restore_preserves_default_and_overrides(self):
        from scripts.invoice_fetch.db_backup import (
            create_verified_database_backup, restore_verified_database_backup,
        )
        path = self.tmp_path / 'backup.db'
        with InvoiceDB(path) as db:
            claim = db.create_claim_group('Backup', reason_detail='组事由')
            inherited = add_invoice(db, 'INHERITED')
            overridden = add_invoice(db, 'OVERRIDDEN')
            assert db.add_invoice_to_claim(claim, inherited)
            assert db.add_invoice_to_claim(claim, overridden)
            assert db.update_invoice_reason(overridden, '独立事由')
        backup_dir = self.tmp_path / 'backups'
        backup = create_verified_database_backup(path, backup_dir=backup_dir)
        with InvoiceDB(path) as db:
            assert db.update_claim_reason(claim, reason_detail='被修改')
            assert db.update_invoice_reason(overridden, None)
        restore_verified_database_backup(backup, path, backup_dir=backup_dir)
        with InvoiceDB(path) as db:
            assert db.get_invoice_reason(inherited) == '组事由'
            assert db.get_invoice_reason(overridden) == '独立事由'
