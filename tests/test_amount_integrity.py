import os
import tempfile
from io import BytesIO
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import openpyxl
from PySide6.QtWidgets import QApplication, QComboBox, QWidget

from scripts.invoice_fetch.claim_export import _normalized_finite_amount
from scripts.invoice_fetch.excel_export import export_excel, _excel_amount_value
from scripts.invoice_fetch.gui.invoice_detail_panel import EditFieldsDialog


class AmountIntegrityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def _dialog(self, amount: str) -> EditFieldsDialog:
        parent = QWidget()
        parent.combo_category = QComboBox(parent)
        parent.combo_category.addItems(["交通", "其他"])
        self.addCleanup(parent.deleteLater)
        dialog = EditFieldsDialog(
            parent,
            "001234567890",
            "2026-09-22",
            amount,
            "交通",
            "Buyer",
            "Seller",
        )
        self.addCleanup(dialog.deleteLater)
        return dialog

    def test_editor_rejects_non_finite_amounts(self):
        for raw in ("NaN", "Infinity", "-Infinity"):
            dialog = self._dialog(raw)
            with patch(
                "scripts.invoice_fetch.gui.invoice_detail_panel.QMessageBox.warning"
            ) as warning:
                self.assertIsNone(dialog.values())
                warning.assert_called_once()

    def test_editor_accepts_thousands_separator_and_normalizes_value(self):
        dialog = self._dialog("1,234.56")
        with patch(
            "scripts.invoice_fetch.gui.invoice_detail_panel.QMessageBox.warning"
        ) as warning:
            values = dialog.values()
        warning.assert_not_called()
        self.assertEqual(values["amount"], "1234.56")

    def test_editor_rejects_malformed_thousands_grouping(self):
        for raw in ("12,34", "1,23,456.78"):
            dialog = self._dialog(raw)
            with patch(
                "scripts.invoice_fetch.gui.invoice_detail_panel.QMessageBox.warning"
            ) as warning:
                self.assertIsNone(dialog.values())
                warning.assert_called_once()

    def test_editor_rejects_subcent_precision(self):
        dialog = self._dialog("0.005")
        with patch(
            "scripts.invoice_fetch.gui.invoice_detail_panel.QMessageBox.warning"
        ) as warning:
            self.assertIsNone(dialog.values())
            warning.assert_called_once()

    def test_claim_amount_boundary_rejects_non_finite_values(self):
        for raw in ("", "NaN", "Infinity", "-Infinity", "not-a-number"):
            with self.assertRaises(ValueError):
                _normalized_finite_amount(raw, invoice_identity="发票号 TEST")
        self.assertEqual(
            _normalized_finite_amount("-12.30", invoice_identity="发票号 TEST"),
            "-12.30",
        )
        self.assertEqual(
            _normalized_finite_amount("1,234.56", invoice_identity="发票号 TEST"),
            "1234.56",
        )
        for raw in ("12,34", "1,23,456.78"):
            with self.assertRaises(ValueError):
                _normalized_finite_amount(raw, invoice_identity="发票号 TEST")
            with self.assertRaises(ValueError):
                _excel_amount_value(raw)
        self.assertEqual(_excel_amount_value("1,234.56"), 1234.56)
        for raw in ("0.005", "1.23001"):
            with self.assertRaises(ValueError):
                _normalized_finite_amount(raw, invoice_identity="发票号 TEST")
            with self.assertRaises(ValueError):
                _excel_amount_value(raw)

    def test_excel_amount_cells_are_numeric_but_invoice_number_stays_text(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "out.xlsx"
            export_excel(
                [
                    {
                        "invoice_number": "00123456789012345678",
                        "invoice_code": "000012345678",
                        "invoice_date": "2026-09-22",
                        "expense_date": "2026-09-22",
                        "amount": "100.25",
                        "total_amount": "128.50",
                        "seller_name": "Seller",
                        "buyer_name": "Buyer",
                        "category": "交通",
                    }
                ],
                path,
            )
            # Read into memory before parsing so openpyxl cannot retain
            # an OS-level handle to the TemporaryDirectory file on Windows.
            workbook_bytes = path.read_bytes()
            wb = openpyxl.load_workbook(BytesIO(workbook_bytes), data_only=False)
            try:
                ws = wb["发票汇总"]
                headers = {cell.value: cell.column for cell in ws[1]}
                number = ws.cell(2, headers["发票号码"])
                amount = ws.cell(2, headers["金额(税前)"])
                total = ws.cell(2, headers["价税合计"])

                self.assertEqual(number.data_type, "s")
                self.assertEqual(number.value, "00123456789012345678")
                self.assertEqual(amount.data_type, "n")
                self.assertEqual(total.data_type, "n")
                self.assertAlmostEqual(amount.value, 100.25, places=2)
                self.assertAlmostEqual(total.value, 128.50, places=2)
                self.assertEqual(amount.number_format, "#,##0.00")
                self.assertEqual(total.number_format, "#,##0.00")
            finally:
                wb.close()


    def test_excel_exception_sheet_amount_cell_is_numeric(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "exception_amounts.xlsx"
            export_excel(
                [
                    {
                        "invoice_number": "1002",
                        "invoice_date": "2026-05-19",
                        "total_amount": "88.00",
                        "category": "打车出行",
                        "missing_extra": 1,
                        "parse_note": "缺少行程记录",
                        "attachment_path": "",
                    },
                ],
                path,
            )
            # Read into memory before parsing so openpyxl cannot retain
            # an OS-level handle to the TemporaryDirectory file on Windows.
            workbook_bytes = path.read_bytes()
            wb = openpyxl.load_workbook(BytesIO(workbook_bytes), data_only=False)
            try:
                ws = wb["异常待处理"]
                headers = {cell.value: cell.column for cell in ws[1]}
                total = ws.cell(2, headers["价税合计"])
                invoice_number = ws.cell(2, headers["发票号码"])

                # A text amount would make SUM()/pivot on this sheet return 0.
                self.assertEqual(total.data_type, "n")
                self.assertIsInstance(total.value, (int, float))
                self.assertAlmostEqual(float(total.value), 88.0, places=2)
                self.assertEqual(total.number_format, "#,##0.00")
                # Non-amount identifier columns keep their text type.
                self.assertEqual(invoice_number.data_type, "s")
                self.assertEqual(invoice_number.value, "1002")
            finally:
                wb.close()

    def test_excel_exception_sheet_empty_amount_stays_empty(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "exception_empty_amount.xlsx"
            export_excel(
                [
                    {
                        "invoice_number": "2001",
                        "invoice_date": "2026-05-20",
                        "total_amount": "",
                        "category": "其他",
                        "missing_extra": 1,
                        "attachment_path": "",
                    },
                ],
                path,
            )
            workbook_bytes = path.read_bytes()
            wb = openpyxl.load_workbook(BytesIO(workbook_bytes), data_only=False)
            try:
                ws = wb["异常待处理"]
                headers = {cell.value: cell.column for cell in ws[1]}
                total = ws.cell(2, headers["价税合计"])

                # Empty amounts must stay empty instead of becoming 0.
                self.assertIn(total.value, (None, ""))
                self.assertNotIsInstance(total.value, (int, float))
            finally:
                wb.close()

    def test_excel_exception_sheet_keeps_formula_injection_protection(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "exception_formulas.xlsx"
            export_excel(
                [
                    {
                        "invoice_number": "3001",
                        "invoice_date": "2026-06-01",
                        "total_amount": "12.30",
                        "category": "其他",
                        "missing_extra": 1,
                        "parse_note": "=SUM(A1:A2)",
                        "mail_subject": "@SUM(A1:A2)",
                        "attachment_path": "",
                    },
                ],
                path,
            )
            workbook_bytes = path.read_bytes()
            wb = openpyxl.load_workbook(BytesIO(workbook_bytes), data_only=False)
            try:
                ws = wb["异常待处理"]
                headers = {cell.value: cell.column for cell in ws[1]}
                note = ws.cell(2, headers["解析备注"])
                subject = ws.cell(2, headers["邮件主题"])
                total = ws.cell(2, headers["价税合计"])

                self.assertEqual(note.value, "'=SUM(A1:A2)")
                self.assertEqual(note.data_type, "s")
                self.assertEqual(subject.value, "'@SUM(A1:A2)")
                self.assertEqual(subject.data_type, "s")
                self.assertEqual(total.data_type, "n")
                self.assertAlmostEqual(float(total.value), 12.30, places=2)
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows():
                        for cell in row:
                            self.assertNotEqual(
                                cell.data_type,
                                "f",
                                f"formula leaked in {sheet.title}!{cell.coordinate}",
                            )
            finally:
                wb.close()

    def test_excel_ledger_sheet_amount_behaviour_is_unchanged(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "ledger_amounts.xlsx"
            export_excel(
                [
                    {
                        "invoice_number": "00123456789012345678",
                        "invoice_date": "2026-09-22",
                        "amount": "100.25",
                        "total_amount": "128.50",
                        "tax_amount": "8.50",
                        "seller_name": "Seller",
                        "buyer_name": "Buyer",
                        "category": "交通",
                        "attachment_path": "attachments/2026-09-22/a.pdf",
                    },
                    {
                        "invoice_number": "4001",
                        "invoice_date": "2026-09-23",
                        "amount": "",
                        "total_amount": "",
                        "tax_amount": "",
                        "seller_name": "Seller",
                        "buyer_name": "Buyer",
                        "category": "交通",
                        "attachment_path": "attachments/2026-09-23/b.pdf",
                    },
                ],
                path,
            )
            workbook_bytes = path.read_bytes()
            wb = openpyxl.load_workbook(BytesIO(workbook_bytes), data_only=False)
            try:
                ws = wb["发票汇总"]
                headers = {cell.value: cell.column for cell in ws[1]}
                for label, expected in (
                    ("金额(税前)", 100.25),
                    ("价税合计", 128.50),
                    ("税额", 8.50),
                ):
                    cell = ws.cell(2, headers[label])
                    self.assertEqual(cell.data_type, "n")
                    self.assertAlmostEqual(float(cell.value), expected, places=2)
                    self.assertEqual(cell.number_format, "#,##0.00")
                    empty = ws.cell(3, headers[label])
                    self.assertIn(empty.value, (None, ""))
                    self.assertNotIsInstance(empty.value, (int, float))
            finally:
                wb.close()


if __name__ == "__main__":
    unittest.main()
